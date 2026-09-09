import sqlite3,json,uuid,hashlib,re,io,zipfile
from datetime import datetime,timezone
from pathlib import Path
from contextlib import contextmanager
from .domain import Persona,PRESETS,Settings
MAX_SIZE=20*1024*1024
MAX_TEXT=2_000_000
class Store:
 def __init__(self,root):
  self.root=root;root.mkdir(parents=True,exist_ok=True)
  for x in ['knowledge','generated','avatars','codex-runtime']:(root/x).mkdir(exist_ok=True)
  with self.db() as c:
   c.executescript('''CREATE TABLE IF NOT EXISTS personas(id TEXT PRIMARY KEY,body TEXT); CREATE TABLE IF NOT EXISTS settings(id INTEGER PRIMARY KEY,body TEXT); CREATE TABLE IF NOT EXISTS docs(id TEXT PRIMARY KEY,name TEXT,size INTEGER,hash TEXT UNIQUE,updated TEXT,chunks INTEGER); CREATE TABLE IF NOT EXISTS chunks(id INTEGER PRIMARY KEY,doc_id TEXT,n INTEGER,text TEXT); CREATE VIRTUAL TABLE IF NOT EXISTS search USING fts5(tokens,title,chunk_id UNINDEXED); CREATE TABLE IF NOT EXISTS conversations(id TEXT PRIMARY KEY,persona TEXT,thread TEXT,history TEXT);''')
   if not c.execute('SELECT 1 FROM personas').fetchone():
    for i,n in enumerate(['林宇辰','陳雅婷','王子豪','張若晴']):
     p=Persona(name=n,avatar=PRESETS[i]['avatar'],scene=['desk','store','product','care'][i]);c.execute('INSERT INTO personas VALUES(?,?)',(uuid.uuid4().hex,p.model_dump_json()))
   c.execute('INSERT OR IGNORE INTO settings VALUES(1,?)',(Settings().model_dump_json(),))
 @contextmanager
 def db(self):
  c=sqlite3.connect(self.root/'app.db');c.row_factory=sqlite3.Row
  try:
   with c:yield c
  finally:c.close()
 def settings(self):
  with self.db() as c:return json.loads(c.execute('SELECT body FROM settings WHERE id=1').fetchone()[0])
 def personas(self):
  with self.db() as c:return [dict(id=r['id'],**json.loads(r['body'])) for r in c.execute('SELECT * FROM personas')]
 def save_persona(self,p,id=None):
  id=id or uuid.uuid4().hex
  with self.db() as c:c.execute('INSERT OR REPLACE INTO personas VALUES(?,?)',(id,p.model_dump_json()))
  return dict(id=id,**p.model_dump())
 def docs(self):
  with self.db() as c:return [dict(r,status='indexed') for r in c.execute('SELECT id,name,size,updated,chunks FROM docs ORDER BY updated DESC')]
 def add_doc(self,name,content,mime):
  if not name or '/' in name or '\\' in name or '..' in name or len(name)>180 or any(ord(x)<32 for x in name):raise ValueError('檔名不安全。')
  if not content or len(content)>MAX_SIZE:raise ValueError('檔案需介於 1 byte 與 20 MB。')
  ext=Path(name).suffix.lower();valid={'.txt':{'text/plain'},'.md':{'text/markdown','text/plain'},'.pdf':{'application/pdf'},'.docx':{'application/vnd.openxmlformats-officedocument.wordprocessingml.document'}}
  if ext not in valid or mime.split(';')[0] not in valid[ext]:raise ValueError('副檔名與 MIME 不符合支援格式。')
  digest=hashlib.sha256(content).hexdigest()
  with self.db() as c:
   r=c.execute('SELECT id FROM docs WHERE hash=?',(digest,)).fetchone()
   if r:return dict(id=r['id'],unchanged=True)
  try:
   if ext in {'.txt','.md'}:
    text=content.decode('utf-8-sig')
    if '\x00' in text or sum(ord(x)<32 and x not in '\n\r\t' for x in text)>0:raise ValueError()
   elif ext=='.pdf':
    if not content.startswith(b'%PDF-'):raise ValueError()
    from pypdf import PdfReader
    pdf=PdfReader(io.BytesIO(content))
    if pdf.is_encrypted or len(pdf.pages)>1000:raise ValueError()
    parts=[];count=0
    for page in pdf.pages:
     t=page.extract_text() or '';count+=len(t)
     if count>MAX_TEXT:raise ValueError()
     parts.append(t)
    text='\n'.join(parts)
   else:
    if not content.startswith(b'PK\x03\x04'):raise ValueError()
    with zipfile.ZipFile(io.BytesIO(content)) as z:
     if sum(i.file_size for i in z.infolist())>40*1024*1024 or 'word/document.xml' not in z.namelist():raise ValueError()
     if any('vbaProject' in i.filename for i in z.infolist()):raise ValueError()
    from docx import Document
    doc=Document(io.BytesIO(content));text='\n'.join([p.text for p in doc.paragraphs]+[' '.join(cell.text for cell in row.cells) for table in doc.tables for row in table.rows])
   if not text.strip() or len(text)>MAX_TEXT:raise ValueError()
  except Exception:raise ValueError('內容格式不符、受保護、過大或沒有可擷取文字；掃描 PDF 請先轉文字。') from None
  id=uuid.uuid4().hex;blocks=[text[i:i+700] for i in range(0,len(text),620)];path=self.root/'knowledge'/id
  path.write_bytes(content)
  try:
   with self.db() as c:
    c.execute('INSERT INTO docs VALUES(?,?,?,?,?,?)',(id,name,len(content),digest,datetime.now(timezone.utc).isoformat(),len(blocks)))
    for i,b in enumerate(blocks):
     cid=c.execute('INSERT INTO chunks(doc_id,n,text) VALUES(?,?,?)',(id,i+1,b)).lastrowid
     c.execute('INSERT INTO search(tokens,title,chunk_id) VALUES(?,?,?)',(' '.join(tokens(b)),' '.join(tokens(name)),cid))
  except Exception:path.unlink(missing_ok=True);raise
  return dict(id=id,unchanged=False)
 def retrieve(self,query,k):
  ts=tokens(query)[:60]
  if not ts:return []
  with self.db() as c:
   rows=c.execute('SELECT ch.doc_id,ch.n,ch.text,d.name,bm25(search,1,4,0) score FROM search JOIN chunks ch ON ch.id=search.chunk_id JOIN docs d ON d.id=ch.doc_id WHERE search MATCH ? ORDER BY score LIMIT ?',(' OR '.join('"'+x+'"' for x in ts),k)).fetchall()
   return [dict(id=r['doc_id'],paragraph=r['n'],text=r['text'],filename=r['name']) for r in rows]
 def delete_doc(self,id):
  with self.db() as c:
   c.execute('DELETE FROM search WHERE chunk_id IN (SELECT id FROM chunks WHERE doc_id=?)',(id,));c.execute('DELETE FROM chunks WHERE doc_id=?',(id,));c.execute('DELETE FROM docs WHERE id=?',(id,))
   # Purge retained assistant history so deleted facts cannot re-enter a new prompt.
   c.execute("UPDATE conversations SET history='[]',thread=NULL")
  (self.root/'knowledge'/id).unlink(missing_ok=True)
def tokens(s):
 runs=re.findall(r'[a-zA-Z0-9]+|[\u3400-\u9fff]+',s.lower());out=[]
 for w in runs:
  out.extend([w[i:i+2] for i in range(len(w)-1)] if '\u3400'<=w[0]<='\u9fff' and len(w)>1 else [w])
 return list(dict.fromkeys(out))
